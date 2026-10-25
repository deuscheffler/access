"""
Convierte las nóminas de Excel (formato del Ministerio de Educación) en
assets/nomina.json, el archivo de datos que viaja dentro de la app.

Uso:
    pip install openpyxl
    python tool/convertir_nominas.py carpeta_con_los_xlsx

Nombre esperado de cada archivo: 1RO_A_MAT.xlsx, 7MO_C_VESP.xlsx,
INICIAL_3_AÑOS_A_MAT.xlsx, etc. (grado, paralelo, jornada). Se toleran
sufijos como "__3_" y prefijos numéricos como "1792812506442_".

Privacidad: la cédula solo se usa para crear un identificador estable del
estudiante (un hash). Ni la cédula ni el correo se guardan en la app.

La clave de firma de los QR se crea la primera vez y se conserva en las
siguientes ejecuciones; si cambiara, los pases ya entregados dejarían de
ser válidos.
"""

import hashlib
import hmac
import json
import re
import secrets
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "assets" / "nomina.json"
PATRON_BASICA = re.compile(r"^(\d+)[A-Z]*_([A-Z])_(MAT|VESP)", re.I)
PATRON_INICIAL = re.compile(r"^INICIAL_(\d)_A\w*OS_([A-Z])_(MAT|VESP)", re.I)
# Código de grado: Inicial 3 años = -2, Inicial 4 años = -1, Básica = 1..10.
GRADO_INICIAL = {3: -2, 4: -1}
PARTICULAS = {"de", "del", "la", "las", "los", "y"}


def titulo(nombre: str) -> str:
    palabras = nombre.strip().lower().split()
    salida = []
    for i, p in enumerate(palabras):
        salida.append(p if (i > 0 and p in PARTICULAS) else p.capitalize())
    return " ".join(salida)


def texto(c) -> str:
    return "" if c is None else str(c).strip()


def leer_archivo(ruta: Path):
    # Quita prefijos numéricos largos que agregan algunos navegadores
    # (ej. 1792812506442_7MO_C_VESP.xlsx).
    nombre = unicodedata.normalize("NFC", ruta.stem)
    nombre = re.sub(r"^\d{6,}_", "", nombre)
    m = PATRON_INICIAL.match(nombre)
    if m and int(m.group(1)) in GRADO_INICIAL:
        grado = GRADO_INICIAL[int(m.group(1))]
    else:
        m = PATRON_BASICA.match(nombre)
        if not m:
            print(f"  Omitido (nombre no reconocido): {ruta.name}")
            return None
        grado = int(m.group(1))
    paralelo, jornada = m.group(2).upper(), m.group(3).upper()

    wb = load_workbook(ruta, data_only=True)
    ws = wb[wb.sheetnames[0]]
    institucion, periodo, filas = "", "", []
    for fila in ws.iter_rows(values_only=True):
        celdas = [texto(c) for c in fila if texto(c)]
        if not celdas:
            continue
        if "ESCUELA" in celdas[0].upper() or "UNIDAD EDUCATIVA" in celdas[0].upper():
            institucion = celdas[0]
        elif re.fullmatch(r"\d{4}\s*-\s*\d{4}", celdas[0]):
            periodo = celdas[0]
        elif re.fullmatch(r"\d+", celdas[0]) and len(celdas) >= 3:
            filas.append((celdas[1], celdas[2]))
    return grado, paralelo, "M" if jornada == "MAT" else "V", institucion, periodo, filas


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    carpeta = Path(sys.argv[1])

    clave = None
    if SALIDA.exists():
        clave = json.loads(SALIDA.read_text(encoding="utf-8")).get("clave")
    clave = clave or secrets.token_hex(32)

    estudiantes, cursos, institucion, periodo = [], {}, "", ""
    for ruta in sorted(carpeta.glob("*.xlsx")):
        datos = leer_archivo(ruta)
        if not datos:
            continue
        grado, paralelo, jornada, inst, per, filas = datos
        institucion = institucion or inst
        periodo = periodo or per
        clave_curso = f"{grado}{paralelo}{jornada}"
        if clave_curso in cursos:
            print(f"  Aviso: {ruta.name} repite el curso {clave_curso}; se omite.")
            continue
        cursos[clave_curso] = len(filas)
        for cedula, nombre in filas:
            ident = hmac.new(clave.encode(), cedula.encode(), hashlib.sha256).hexdigest()[:10]
            estudiantes.append({
                "id": ident,
                "n": titulo(nombre),
                "g": grado,
                "p": paralelo,
                "j": jornada,
            })

    nombre_inst = re.sub(r"\s*-\s*\w*\d+\w*\s*$", "", institucion).strip()
    nombre_inst = unicodedata.normalize("NFC", nombre_inst.replace("“", "«").replace("”", "»"))

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    SALIDA.write_text(json.dumps({
        "institucion": titulo(nombre_inst).replace("«república", "«República"),
        "periodo": periodo,
        "clave": clave,
        "estudiantes": estudiantes,
    }, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    print(f"\n{len(estudiantes)} estudiantes en {len(cursos)} cursos -> {SALIDA}")
    for k in sorted(cursos, key=lambda c: (c[-1], int(re.match(r"-?\d+", c).group()), c)):
        print(f"  {k}: {cursos[k]}")


if __name__ == "__main__":
    main()
