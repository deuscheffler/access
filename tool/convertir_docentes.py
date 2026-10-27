"""
Convierte la lista de docentes (Excel con hojas MATUTINA y VESPERTINA) en
assets/docentes.json, que la app usa para enviar los pases de cada paralelo
al WhatsApp de su docente.

Uso:
    python tool/convertir_docentes.py DATOS_PARA_APLICACION_QR.xlsx
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "assets" / "docentes.json"
GRADOS = {
    "INICIAL 3 ANOS": -2, "INICIAL 4 ANOS": -1, "PRIMERO": 1, "SEGUNDO": 2,
    "TERCERO": 3, "CUARTO": 4, "QUINTO": 5, "SEXTO": 6, "SEPTIMO": 7,
    "OCTAVO": 8, "NOVENO": 9, "DECIMO": 10,
}
PARTICULAS = {"de", "del", "la", "las", "los", "y"}


def sin_tildes(s: str) -> str:
    s = unicodedata.normalize("NFD", s)
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def titulo(nombre: str) -> str:
    p = nombre.strip().lower().split()
    return " ".join(w if (i and w in PARTICULAS) else w.capitalize() for i, w in enumerate(p))


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    wb = load_workbook(sys.argv[1], data_only=True)
    docentes, escuela = {}, {}
    for ws in wb.worksheets:
        jornada = "V" if "VESP" in ws.title.upper() else "M"
        for fila in ws.iter_rows(values_only=True):
            celdas = ["" if c is None else str(c).strip() for c in fila]
            if celdas and "TELEF" in celdas[0].upper() and not escuela:
                t = celdas[0]
                m = re.search(r"(PASAJE[^–\-\n]*)", t, re.I)
                escuela["direccion"] = titulo(m.group(1)).replace("S/n", "S/N") if m else ""
                m = re.search(r"TELEF:\s*([\d ]+)", t, re.I)
                escuela["telefono"] = m.group(1).strip() if m else ""
                m = re.search(r"[\w.\-]+@[\w.\-]+", t)
                escuela["correo"] = m.group(0) if m else ""
            if len(celdas) < 5 or not re.fullmatch(r"\d+", celdas[0]):
                continue
            nombre, grado_txt, paralelo, celular = celdas[1:5]
            grado = GRADOS.get(re.sub(r"\s+", " ", sin_tildes(grado_txt).upper()).strip())
            if grado is None:
                print(f"  Grado no reconocido: {grado_txt!r}")
                continue
            clave = f"{grado}{paralelo.upper()}{jornada}"
            docentes[clave] = {
                "nombre": titulo(nombre),
                "celular": re.sub(r"\D", "", celular),
            }
    SALIDA.write_text(json.dumps({"escuela": escuela, "docentes": docentes},
                                 ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(docentes)} docentes -> {SALIDA}")
    print("Escuela:", escuela)


if __name__ == "__main__":
    main()
